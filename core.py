# -*- coding: utf-8 -*-
"""
PDPC administrative interpretation updater -- shared core logic.

This module has NO GUI dependency, so it can be reused by both:
  - app.py            (the desktop window that gets packaged into an executable)
  - a future Google Colab notebook
When switching delivery modes you only swap the outer shell (how files come in /
out and how the run is triggered); this file stays unchanged.

Public functions:
  scrape_all(progress=None) -> list[dict]
      Scrape every administrative interpretation from the official site.
  sync_excel(excel, rows, dry_run=False, progress=None) -> dict
      Compare against an existing Excel and (unless dry_run) append new rows.

Parsing note:
  The official pages embed their data oddly and break HTML tree parsers
  (BeautifulSoup/lxml drop the content section from the main tree). The markup
  for each field is however very regular, so we parse it with regexes directly,
  which is robust and keeps the dependency list small (requests + openpyxl only).
"""

import io
import json
import os
import re
import shutil
import ssl
import time
import unicodedata
from datetime import datetime
from html import unescape
from urllib.parse import urljoin, urlparse

import requests


# ---------------------------------------------------------------------------
# TLS trust: prefer the OS certificate store (applied on import)
# ---------------------------------------------------------------------------
# requests validates server certificates against certifi's own bundled CA list,
# which knows nothing about the private root CA that a company's TLS-inspecting
# proxy re-signs outbound traffic with. On such a network the browser opens
# pdpc.gov.tw fine (IT pushed that root into the Windows certificate store) while
# this tool dies with "self-signed certificate in certificate chain" -- so point
# Python at the OS store too, and whatever IT installed is trusted automatically.
#
# Needs CPython >= 3.10 and the truststore package; a miss here is NOT fatal,
# _get() then degrades to an unverified retry.
TRUSTSTORE_ACTIVE = False


def _install_os_truststore():
    global TRUSTSTORE_ACTIVE
    try:
        import truststore
        truststore.inject_into_ssl()
        TRUSTSTORE_ACTIVE = True
    except Exception:  # noqa: BLE001 -- package missing / old Python / patch failed
        pass


_install_os_truststore()


# ---------------------------------------------------------------------------
# openpyxl whitespace fix (applied on import)
# ---------------------------------------------------------------------------
# openpyxl's whitespace() helper only marks a <t> run with xml:space="preserve"
# when the run has NON-whitespace content (its guard is `if stripped and ...`).
# A whitespace-ONLY run -- e.g. a bare "\n" line-break inside a rich-text cell --
# is therefore written WITHOUT xml:space="preserve", so Excel discards the run
# and then reports the whole workbook as corrupt / unopenable. This bites both
# the red rich-text we write AND any pre-existing rich-text cells the user
# formatted by hand, the moment the workbook is re-saved.
#
# We patch the helper to preserve any run whose text differs from its stripped
# form (leading/trailing OR wholly whitespace), which is the correct OOXML rule.
# whitespace() is imported by-name into several openpyxl modules, so we rebind
# the name in each of them.
def _install_whitespace_preserve_fix():
    from openpyxl.xml.constants import XML_NS

    def whitespace(node):
        if node.text is not None and node.text != node.text.strip():
            node.set("{%s}space" % XML_NS, "preserve")

    import openpyxl.xml.functions as _functions
    _functions.whitespace = whitespace
    for modname in ("openpyxl.cell.rich_text", "openpyxl.cell._writer",
                    "openpyxl.descriptors.nested"):
        try:
            mod = __import__(modname, fromlist=["whitespace"])
            if hasattr(mod, "whitespace"):
                mod.whitespace = whitespace
        except Exception:  # noqa: BLE001 -- never let a patch failure break import
            pass


_install_whitespace_preserve_fix()


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
BASE_URL = "https://www.pdpc.gov.tw"
LIST_URL = "https://www.pdpc.gov.tw/News_Html/100/"
# The only host we will fetch from or hand back as a clickable link. The site's
# own JSON supplies absolute URLs, so without this an edited list page (exactly
# what the TLS fallback below tolerates) could point us at an intranet host.
ALLOWED_HOST = "pdpc.gov.tw"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    )
}
REQUEST_DELAY = 0.4          # polite delay between requests (seconds)
MAX_RETRY = 3
TIMEOUT = 30

# Keywords used to auto-detect Excel columns from the header row.
# The master sheet has three columns: 條號 (article), 條文內容 (article text, left
# untouched) and 最新函釋 -- the latter packs every interpretation for that article
# (one "date+doc_no" line each).
# Order matters: fields are tried per header cell in dict order, so funhao's
# more-specific keywords go FIRST -- otherwise article's bare "條" fallback
# would also claim a 函號 header like "最新函釋(共147條)".
COL_KEYWORDS = {
    "funhao": ["最新函釋", "函釋", "函號"],
    "article": ["條號", "條文", "條次", "條"],
}

# "第2條" / "第1-1條" style article key, used to match a scraped article to an
# Excel row regardless of any prefix like "個人資料保護法 ".
_ARTICLE_KEY_RE = re.compile(r"第[0-9]+(?:-[0-9]+)?條")
# Leading 民國 date on a 函號 line. The site uses two formats -- "115.1.23"
# and "114年1月24日" -- and a 發文字號 always begins with a Chinese agency name,
# so stripping any leading run of digits / date punctuation / 年月日 is safe.
_DATE_PREFIX_RE = re.compile(r"^[\s\d./\-年月日]+")
# Parse a leading 民國 date for sorting; handles "115.1.23" and "114年1月24日".
_DATE_PARSE_RE = re.compile(r"^\s*(\d+)\s*[.年]\s*(\d+)\s*[.月]\s*(\d+)")
# The 字第<數字>號 serial inside a 發文字號. The same interpretation is published
# with either the full agency 全銜 or a 簡稱 prefix ("個人資料保護委員會籌備處個資
# 籌法字第..." vs "個資籌法字第..."), but the serial after 字第 is identical, so it
# is the stable part to compare on.
_DOC_NO_KEY_RE = re.compile(r"字第(\d+)號")
# The 機關+字軌 prefix of a 發文字號 = everything up to and including the 字 that
# precedes 第<序號>號 (e.g. "個資籌法字", "國家發展委員會發法字").
_DOC_PREFIX_RE = re.compile(r"(.*?)字第\d+號\s*$")

# The PDPC site emits some 發文字號 in an abbreviated 字軌 form, dropping the
# issuing agency 全銜 (it lists "個資籌法字第..." rather than the full "個人資料保護
# 委員會籌備處個資籌法字第..."). The master, however, records the full 機關名稱.
# This built-in map expands the known abbreviations so every line written/shown
# carries the full 全銜 like the master's existing entries. It is supplemented at
# run time by mappings derived from the master's own entries (see
# build_agency_expansion), so agencies the user has already entered in 全銜 form
# expand automatically without editing this table.
AGENCY_FULL_PREFIX = {
    "個資籌法字": "個人資料保護委員會籌備處個資籌法字",
    "發法字": "國家發展委員會發法字",
}

# ARGB colour used to highlight lines this run newly added to a 函號 cell, so the
# user can spot exactly what changed. Red.
NEW_LINE_COLOR = "FFFF0000"
# Prefix used when adding a brand-new article row. The master sheet's 條號 column
# stores bare keys like "第2條", so no prefix is added.
ARTICLE_LABEL_PREFIX = ""

# One regex match per labelled field on an interpretation page. Each field
# lives in its own <li>, so we capture up to </li>; this also correctly joins
# multi-paragraph 要旨 (e.g. "<p></p><p>...</p><p>...</p>") that a single
# non-greedy <p>...</p> would truncate to an empty leading paragraph.
_FIELD_RE = re.compile(
    r"<strong>\s*(發文字號|發文日期|要旨)\s*[:：]\s*</strong>\s*(.*?)</li>",
    re.S,
)
_HREF_RE = re.compile(r'href="([^"]+)"')
_TAG_RE = re.compile(r"<[^>]+>")


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def _log(progress, msg):
    """Forward a progress message to the shell (GUI / Colab). No-op if None."""
    if progress:
        try:
            progress(msg)
        except Exception:
            pass


def normalize_doc_no(text):
    """
    Normalize a document number so it can be used as the comparison key.
      - NFKC folds full-width digits/letters/brackets to half-width
      - strip all whitespace
    """
    if text is None:
        return ""
    s = unicodedata.normalize("NFKC", str(text))
    return re.sub(r"\s+", "", s).strip()


def _clean_text(fragment):
    """Strip HTML tags and unescape entities from an inner-HTML fragment."""
    text = _TAG_RE.sub("", fragment)
    text = unescape(text)
    return re.sub(r"\s+", " ", text).strip()


# Set once certificate verification has been abandoned; stays set for the rest
# of the process so we do not re-fail the handshake on every one of the ~30 pages.
_INSECURE_TLS = False


def _safe_link(href, fallback):
    """Resolve href against the site, falling back if it points off-site."""
    if not href:
        return fallback
    url = urljoin(BASE_URL, href)
    return url if is_allowed_url(url) else fallback


class FetchRefused(RuntimeError):
    """A request we refuse to make or to trust -- never worth retrying."""


def is_allowed_url(url):
    """True if url points at the official site (or a subdomain of it)."""
    host = (urlparse(str(url or "")).hostname or "").lower()
    return host == ALLOWED_HOST or host.endswith("." + ALLOWED_HOST)


def _is_cert_verification_failure(exc):
    """
    True only if exc was caused by the certificate chain failing validation.

    Matching on the message text is NOT reliable: with truststore active the OS
    reports the failure itself, and macOS returns a LOCALISED string with no
    "CERTIFICATE_VERIFY_FAILED" in it at all (observed: 「*.badssl.com」憑證不受
    信任). Both paths do raise ssl.SSLCertVerificationError, so walk the wrapper
    chain -- requests.SSLError -> urllib3.MaxRetryError -> urllib3.SSLError ->
    ssl.SSLCertVerificationError -- and match on the type, keeping the text
    check only as a backstop.
    """
    seen, cur = set(), exc
    while isinstance(cur, BaseException) and id(cur) not in seen:
        seen.add(id(cur))
        if isinstance(cur, ssl.SSLCertVerificationError):
            return True
        if "CERTIFICATE_VERIFY_FAILED" in str(cur):
            return True
        nxt = getattr(cur, "reason", None)          # urllib3's MaxRetryError
        if not isinstance(nxt, BaseException):
            nxt = cur.__cause__ or cur.__context__
        if not isinstance(nxt, BaseException) and cur.args:
            nxt = cur.args[0] if isinstance(cur.args[0], BaseException) else None
        cur = nxt
    return False


def _silence_insecure_warnings():
    """Mute urllib3's per-request InsecureRequestWarning once we stop verifying."""
    try:
        import urllib3
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
    except Exception:  # noqa: BLE001
        pass


def _get(url, session, progress=None):
    """
    GET with retry/backoff; return decoded HTML text.

    On the FIRST certificate-verification failure we give up verifying for the
    rest of the run rather than failing outright. Rationale: behind a company
    proxy that inspects TLS, verification cannot succeed unless the private root
    CA is trusted (which _install_os_truststore() already tried), and the
    alternative is a tool that flatly does not work there. What we lose is the
    guarantee that the peer really is pdpc.gov.tw -- acceptable here because we
    only READ public pages and send no credentials, so the worst case is a
    tampered interpretation text, not leaked data.
    """
    global _INSECURE_TLS
    if not is_allowed_url(url):
        raise FetchRefused(f"Refusing to fetch a non-{ALLOWED_HOST} URL: {url}")
    last_err = None
    for attempt in range(1, MAX_RETRY + 1):
        try:
            resp = session.get(url, headers=HEADERS, timeout=TIMEOUT,
                               verify=not _INSECURE_TLS)
            # Redirects are followed, so re-check where we actually landed and
            # drop the body unparsed if it left the official site.
            if not is_allowed_url(resp.url):
                raise FetchRefused(
                    f"Refusing a redirect off {ALLOWED_HOST}: {url} -> {resp.url}")
            resp.raise_for_status()
            # Trust a charset declared in the Content-Type header; only guess
            # from the body (apparent_encoding, slow and fallible) without one.
            # requests defaults undeclared text/* to ISO-8859-1, which would
            # mangle the Chinese pages.
            content_type = resp.headers.get("Content-Type", "")
            if "charset" not in content_type.lower():
                resp.encoding = resp.apparent_encoding or "utf-8"
            return resp.text
        except FetchRefused:
            raise                       # our own refusal; retrying cannot help
        except requests.exceptions.SSLError as e:
            last_err = e
            # Degrade ONLY on a real chain-validation failure. SSLError also
            # covers handshake resets, protocol/version errors and mid-handshake
            # EOF -- routine blips on a corporate network -- and treating those
            # as proof of interception would drop verification for the rest of
            # the run over a hiccup, while telling the user something false.
            if _is_cert_verification_failure(e) and not _INSECURE_TLS:
                _INSECURE_TLS = True
                _silence_insecure_warnings()
                _log(progress, "⚠ 無法驗證網站憑證（可能是公司網路的 SSL 攔截，"
                               "也可能是連線被干擾）；已略過憑證驗證繼續執行"
                               "（本工具只讀取公開網頁，不會外送任何資料）。"
                               "若不確定原因，請向 IT 確認。")
                continue  # retry this same URL unverified, no backoff needed
            time.sleep(REQUEST_DELAY * attempt)
        except Exception as e:  # noqa: BLE001
            last_err = e
            time.sleep(REQUEST_DELAY * attempt)
    raise RuntimeError(f"Failed to fetch page after {MAX_RETRY} retries: {url}\n{last_err}")


# ---------------------------------------------------------------------------
# Step 1: find every article that has an administrative-interpretation link
# ---------------------------------------------------------------------------
def parse_article_list(html):
    """
    Parse the DefaultData JSON embedded in the News_Html/100 list page and
    return every article that carries an administrative-interpretation link,
    as [(article_no, interpretation_url), ...].

    The count is discovered dynamically -- it is NOT hard-coded, so links added
    in future site updates are picked up automatically.
    """
    m = re.search(r"let\s+DefaultData\s*=\s*(\[.*?\]);", html, flags=re.S)
    if not m:
        raise RuntimeError(
            "Could not find DefaultData on the list page; the site layout may have changed."
        )
    data = json.loads(m.group(1))

    results = []
    seen = set()
    for entry in data:
        art_no = (entry.get("條號") or "").strip()
        links_html = entry.get("LinksHtml") or ""
        info = entry.get("相關資訊") or ""

        href = _find_interpretation_link(links_html, info)
        if not href:
            continue
        url = urljoin(BASE_URL, href)
        if not is_allowed_url(url):     # absolute link pointing off-site
            continue
        if url in seen:
            continue
        seen.add(url)
        results.append((art_no, url))
    return results


def _find_interpretation_link(links_html, info):
    """Return the href of the administrative-interpretation link, or None."""
    # Primary: the "相關資訊" string contains 【行政函釋】(url)
    m = re.search(r"【行政函釋】\s*\((https?://[^)]+)\)", info)
    if m:
        return m.group(1)

    # Fallback: scan anchors in LinksHtml for one labelled 行政函釋
    for a_tag, inner in re.findall(r"(<a\b[^>]*>)(.*?)</a>", links_html, flags=re.S):
        if "行政函釋" in a_tag or "行政函釋" in inner:
            href = _HREF_RE.search(a_tag)
            if href:
                return href.group(1)
    return None


# ---------------------------------------------------------------------------
# Step 2: parse a single interpretation page into individual records
# ---------------------------------------------------------------------------
def parse_interpretation_page(html, art_no, source_url=""):
    """
    Parse a News_Content/101/{id} page and return every interpretation under
    that article as dicts:
      {article, doc_no, doc_no_norm, doc_date, summary, detail_url, article_url}

    Each record is a (發文字號, 發文日期, 要旨) triple in document order; a new
    record starts at every 發文字號, tolerating missing fields.

    detail_url is the per-interpretation deep link if the 發文字號 markup carries
    one; otherwise it falls back to source_url (the page this interpretation was
    scraped from), so a usable link is always present when source_url is given.
    article_url is always source_url -- the article-level interpretation page
    (one level up from the per-interpretation deep link).
    """
    rows = []
    current = None

    for label, fragment in _FIELD_RE.findall(html):
        if label == "發文字號":
            if current and current["doc_no"]:
                rows.append(current)
            value = _clean_text(fragment)
            href = _HREF_RE.search(fragment)
            current = {
                "article": art_no,
                "doc_no": value,
                "doc_no_norm": normalize_doc_no(value),
                "doc_date": "",
                "summary": "",
                # Shown to the user as a clickable 連結, so it must stay on-site
                # even if the page was tampered with; fall back to the page URL.
                "detail_url": _safe_link(href.group(1) if href else None, source_url),
                "article_url": source_url,
            }
        elif current is not None:
            if label == "發文日期":
                current["doc_date"] = _clean_text(fragment)
            elif label == "要旨":
                current["summary"] = _clean_text(fragment)

    if current and current["doc_no"]:
        rows.append(current)
    return rows


# ---------------------------------------------------------------------------
# Public: scrape everything
# ---------------------------------------------------------------------------
def scrape_all(progress=None):
    """Scrape every administrative interpretation; return list[dict]."""
    # Per RUN, not per process: the GUI is long-lived, so a dry-run that had to
    # drop verification must not silently carry that into the update that
    # actually writes. Re-degrading simply warns again, which is the point.
    global _INSECURE_TLS
    _INSECURE_TLS = False

    session = requests.Session()
    _log(progress, "Loading the article list page...")
    list_html = _get(LIST_URL, session, progress)
    articles = parse_article_list(list_html)
    _log(progress, f"Found {len(articles)} articles with interpretation links.")

    all_rows = []
    for idx, (art_no, url) in enumerate(articles, 1):
        _log(progress, f"[{idx}/{len(articles)}] Parsing interpretations for {art_no}...")
        try:
            html = _get(url, session, progress)
            all_rows.extend(parse_interpretation_page(html, art_no, source_url=url))
        except Exception as e:  # noqa: BLE001
            _log(progress, f"  ! Skipped {art_no} (parse failed): {e}")
        time.sleep(REQUEST_DELAY)

    _log(progress, f"Scrape complete: {len(all_rows)} interpretations total.")
    return all_rows


# ---------------------------------------------------------------------------
# Step 3: Excel column detection, comparison and write-back
# ---------------------------------------------------------------------------
def _detect_columns(ws):
    """
    Auto-detect columns from the header row (row 1) -> {field: column_index}.
    Detection succeeds only if both the article and 函號 columns are found.
    """
    header_cells = next(ws.iter_rows(min_row=1, max_row=1), [])
    headers = {}
    for cell in header_cells:
        text = unicodedata.normalize("NFKC", str(cell.value or "")).strip()
        if not text:
            continue
        for field, keywords in COL_KEYWORDS.items():
            if field in headers:
                continue
            if any(kw.lower() in text.lower() for kw in keywords):
                headers[field] = cell.column
                break
    return headers


def article_key(label):
    """Reduce an article label to its '第X條' key for matching across rows."""
    s = unicodedata.normalize("NFKC", str(label or ""))
    m = _ARTICLE_KEY_RE.search(s)
    return m.group(0) if m else s.strip()


def doc_no_from_line(line):
    """Extract the normalized 發文字號 from a '日期+字號' 函號 line."""
    s = unicodedata.normalize("NFKC", str(line or "")).strip()
    if not s:
        return ""
    return normalize_doc_no(_DATE_PREFIX_RE.sub("", s))


def doc_no_key(text):
    """The 字第<數字>號 serial of a 發文字號, used as the 字號 half of the identity
    key. It is stable across the agency 全銜/簡稱 prefix. Falls back to the fully
    normalized text when no serial is present (so odd inputs still compare)."""
    s = normalize_doc_no(text)
    m = _DOC_NO_KEY_RE.search(s)
    return m.group(1) if m else s


def doc_prefix(text):
    """The 機關+字軌 prefix of a 發文字號 ('個資籌法字', '國家發展委員會發法字', ...),
    i.e. everything up to the 字 before 第<序號>號. Any leading 民國 date is
    stripped first; returns '' if the text has no 字第<序號>號 serial."""
    s = normalize_doc_no(_DATE_PREFIX_RE.sub("", str(text or "")))
    m = _DOC_PREFIX_RE.match(s)
    return m.group(1) + "字" if m else ""


def build_agency_expansion(prefixes):
    """From the set of 機關+字軌 prefixes already present in the master, derive a
    {abbrev: full} expansion map: an abbreviated prefix P maps to the longest
    full prefix F (F != P) that ends with P. This lets agencies the user has
    already entered in 全銜 form expand future abbreviated lines automatically,
    on top of the built-in AGENCY_FULL_PREFIX table."""
    pset = sorted({p for p in prefixes if p})
    mapping = {}
    for p in pset:
        longer = [f for f in pset if f != p and f.endswith(p)]
        if longer:
            mapping[p] = max(longer, key=len)
    return mapping


def expand_agency_prefix(line, extra_map=None):
    """Expand an abbreviated 機關字軌 in a 函號 line to its full 全銜, leaving the
    leading 民國 date, the 字軌 and the 序號 untouched:
      '113.5.28個資籌法字第1130000459號'
        -> '113.5.28個人資料保護委員會籌備處個資籌法字第1130000459號'.
    Already-full lines are returned unchanged (their prefix is not a known
    abbreviation), so this is idempotent. extra_map supplements the built-in
    AGENCY_FULL_PREFIX (e.g. a map derived from the master's own entries)."""
    mapping = AGENCY_FULL_PREFIX if not extra_map else {**AGENCY_FULL_PREFIX, **extra_map}
    s = str(line or "")
    rest = _DATE_PREFIX_RE.sub("", s)          # the 發文字號 part (date stripped)
    date = s[: len(s) - len(rest)]
    m = _DOC_PREFIX_RE.match(rest)
    if not m:
        return s
    prefix = m.group(1) + "字"                  # 機關+字軌, e.g. "個資籌法字"
    full = mapping.get(prefix)
    if not full or rest.startswith(full):
        return s
    return f"{date}{full}{rest[len(prefix):]}"


def to_dotted_date(date_str):
    """Normalize a 發文日期 to dotted form with NO leading zeros:
    '114年1月24日' -> '114.1.24', '114.09.22' -> '114.9.22'.
    Accepts either 年月日 or already-dotted input (zero-padded or not); anything
    unrecognized is returned unchanged. Stripping the zeros makes '114.09.22' and
    '114.9.22' write out identically (one canonical form)."""
    s = unicodedata.normalize("NFKC", (date_str or "").strip())
    m = re.fullmatch(r"\s*(\d+)\s*[.年]\s*(\d+)\s*[.月]\s*(\d+)\s*日?\s*", s)
    return f"{int(m.group(1))}.{int(m.group(2))}.{int(m.group(3))}" if m else s


def normalize_funhao_line(line):
    """Rewrite a 函號 line's leading 民國 date to dotted, zero-stripped form,
    leaving the 發文字號 text untouched:
      '114.09.22個資籌法字第1140001075號' -> '114.9.22個資籌法字第1140001075號'.
    A line with no leading date is returned trimmed but otherwise unchanged."""
    s = str(line or "").strip()
    rest = _DATE_PREFIX_RE.sub("", s)
    date = s[: len(s) - len(rest)]
    return f"{to_dotted_date(date)}{rest}" if date.strip() else rest


def funhao_line(row):
    """Format one scraped interpretation as a 函號 line: '日期字號'
    (date normalized to zero-stripped dotted form so everything written is x.x.x)."""
    date = to_dotted_date(row.get("doc_date"))
    return expand_agency_prefix(f"{date}{row.get('doc_no', '').strip()}")


def parse_roc_date(line):
    """
    Extract the leading 民國 date from a 函號 line as a sortable (y, m, d) tuple.
    Handles "115.1.23" and "114年1月24日". Returns None if no date is found,
    so callers can sort undated lines to the top (treated as oldest).
    """
    s = unicodedata.normalize("NFKC", str(line or ""))
    m = _DATE_PARSE_RE.match(s)
    if not m:
        return None
    return tuple(int(g) for g in m.groups())


def funhao_key(line):
    """Identity key for a 函號 line = the 發文字號序號 only (the 字第<數字>號 serial).
    The date is NOT part of the key: an interpretation is identified solely by its
    字號, so the same 字號 is treated as one entry even if the master's recorded
    date differs from the site's, and regardless of 全銜/簡稱 or date format."""
    return doc_no_key(line)


# Excel evaluates a cell whose text starts with any of these as a formula, and
# openpyxl faithfully types such a string as one (data_type 'f'). Everything we
# write is ultimately site-controlled, so a 條號 of '=HYPERLINK("http://x",…)'
# would land in the master sheet as a LIVE formula -- i.e. whoever controls the
# page (including the proxy the TLS fallback tolerates) could plant WEBSERVICE /
# HYPERLINK / DDE payloads. Forcing the string data type keeps the exact text
# visible while stopping Excel from ever executing it.
_FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")


def _assign_text(cell, value):
    """Assign a site-derived value, forcing Excel to store it as text."""
    cell.value = value
    if isinstance(value, str) and value.startswith(_FORMULA_LEAD):
        cell.data_type = "s"
    return cell


def load_workbook_from(excel):
    """
    Load an openpyxl workbook from a path string or a file-like object.
    (Lets the desktop path and a Colab uploaded file share one entry point.)
    Returns (workbook, original_bytes).

    rich_text=True so that red highlights written on a previous run survive a
    re-run: without it openpyxl reads rich-text cells as plain strings and would
    drop the colour from untouched rows on save. str()-ing a cell value still
    yields plain text, so the comparison logic is unaffected.
    """
    from openpyxl import load_workbook

    if hasattr(excel, "read"):              # file-like / BytesIO
        raw = excel.read()
    else:                                    # path string
        with open(excel, "rb") as f:
            raw = f.read()
    wb = load_workbook(io.BytesIO(raw), rich_text=True)
    return wb, raw


def is_excel_locked(path):
    """
    Best-effort check for whether an .xlsx is currently open elsewhere
    (typically in Excel), so the shell can ask the user to close it before a
    write-back. Returns True if the file looks open/locked.

    No single signal is reliable across platforms, so we use two:
      - Excel drops a hidden owner file "~$<name>" beside the workbook while it
        is open, on BOTH Windows and macOS. This is the only signal on macOS,
        which has no mandatory file lock.
      - Opening the file for writing fails with PermissionError when Excel holds
        an exclusive lock (the Windows case).

    Reading (dry-run) is unaffected on either platform, so this only matters for
    the write-back step. Note: a stale "~$" file left by an Excel crash can cause
    a false positive; the user can delete it. File-like inputs (Colab) can't be
    locked, so they always return False.
    """
    if hasattr(path, "read"):          # file-like (e.g. Colab upload)
        return False
    try:
        folder, name = os.path.split(path)
        if os.path.exists(os.path.join(folder, "~$" + name)):
            return True
    except OSError:
        pass
    try:
        with open(path, "r+b"):        # request write access; harmless if granted
            return False
    except PermissionError:
        return True
    except OSError:
        return False
    return False


def sync_excel(excel, rows, dry_run=False, progress=None):
    """
    Compare scraped interpretations against the existing master Excel and,
    unless dry_run, append the new ones.

    The master sheet has one row per article; its 函號 cell packs every
    interpretation for that article as newline-separated "date+doc_no" lines.
    Comparison is therefore PER ARTICLE, keyed by normalized 發文字號:
      - new interpretation under an existing article -> append a line to its 函號 cell
      - article not yet present                      -> add a new row at the bottom

    Args:
        excel    : Excel path string, or a file-like object (has .read()).
                   Write-back requires a path.
        rows     : output of scrape_all().
        dry_run  : if True, only compare; do not write.

    Returns dict:
        {
          "columns":        detected column mapping,
          "existing_count": interpretation lines already in the Excel,
          "scraped_count":  interpretations scraped from the site,
          "new_rows":       the new interpretations (list[dict], each has article/doc_no/doc_date),
          "new_count":      number of new interpretations,
          "new_articles":   article keys that will get a brand-new row,
          "backup_path":    backup file path (or None),
          "written":        whether the file was actually written,
        }
    """
    from openpyxl.styles import Alignment

    wb, _raw = load_workbook_from(excel)
    ws = wb.active

    cols = _detect_columns(ws)
    if "article" not in cols or "funhao" not in cols:
        raise RuntimeError(
            "Could not detect the required columns in the Excel header row.\n"
            "Row 1 must contain a 條號 column and a 最新函釋 column.\n"
            f"Detected columns: {cols or '(none)'}"
        )
    art_col, fun_col = cols["article"], cols["funhao"]

    # Index existing rows by article key. doc_nos is the set of 字號序號 identity
    # keys, unioned across every row of the same article (for correct comparison);
    # row/lines point at the FIRST row of that article, which is the one new lines
    # get merged into and re-sorted.
    existing = {}
    existing_count = 0
    existing_keys = set()      # distinct 字號 across the whole sheet (不重複計算)
    existing_prefixes = set()  # 機關+字軌 prefixes seen, for 全銜 expansion
    for ridx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        key = article_key(row[art_col - 1].value)
        if not key:
            continue
        raw_lines = [ln for ln in str(row[fun_col - 1].value or "").split("\n") if ln.strip()]
        doc_nos = {funhao_key(ln) for ln in raw_lines}
        existing_prefixes |= {doc_prefix(ln) for ln in raw_lines}
        existing_count += len(doc_nos)
        existing_keys |= doc_nos
        if key in existing:
            existing[key]["doc_nos"] |= doc_nos
        else:
            existing[key] = {"row": ridx, "doc_nos": doc_nos, "lines": raw_lines}
    _log(progress, f"Excel already contains {existing_count} interpretations "
                   f"across {len(existing)} articles.")

    # Group scraped rows by article, find the ones missing per article.
    new_rows = []
    appended = {}          # article key -> [scraped rows to add]
    new_articles = []      # article keys needing a brand-new row
    seen = {}              # article key -> set of identity keys seen this batch
    for r in rows:
        key = article_key(r["article"])
        if not key or not r.get("doc_no", "").strip():
            continue
        dn = funhao_key(funhao_line(r))     # 字號序號, same key as existing
        batch_seen = seen.setdefault(key, set())
        existing_set = existing[key]["doc_nos"] if key in existing else set()
        if dn in existing_set or dn in batch_seen:
            continue
        batch_seen.add(dn)
        new_rows.append(r)
        appended.setdefault(key, []).append(r)
        if key not in existing and key not in new_articles:
            new_articles.append(key)

    # Projected distinct 函釋 total after this run (現有 ∪ 本次新增), so the report
    # / dry-run can show the 共N條 the header will become.
    funhao_total = len(existing_keys | {funhao_key(funhao_line(r)) for r in new_rows})

    result = {
        "columns": cols,
        "existing_count": existing_count,
        "scraped_count": len(rows),
        "new_rows": new_rows,
        "new_count": len(new_rows),
        "new_articles": new_articles,
        "funhao_total": funhao_total,
        "backup_path": None,
        "written": False,
    }
    _log(progress, f"Comparison done: {len(new_rows)} new interpretation(s)"
                   f"{f', {len(new_articles)} new article row(s)' if new_articles else ''}.")

    if dry_run or not new_rows:
        return result

    if hasattr(excel, "read"):
        raise RuntimeError("Write-back needs a file path, not a file-like object.")

    # openpyxl rebuilds the whole workbook on save and models only the parts it
    # understands: macros, charts, images, pivot tables, form controls and
    # custom XML are silently dropped. That is destructive for .xlsm/.xltx, and
    # the extension would still claim otherwise -- so refuse rather than eat them.
    if not str(excel).lower().endswith(".xlsx"):
        raise RuntimeError(
            "只支援 .xlsx 主表。\n\n"
            "本工具寫回時會重建整個活頁簿，巨集（.xlsm）、圖表、圖片、樞紐分析表等\n"
            "會在存檔時遺失，因此不開放寫入。請另存成 .xlsx 後再更新。")

    # Safety net (covers the case where the file is opened DURING the scrape):
    # bail out before backing up, so a locked file leaves no stray .bak behind.
    if is_excel_locked(excel):
        raise RuntimeError("Excel 檔案目前開啟中，無法寫回。請先關閉該檔後再試一次。")

    # Back up the original before writing.
    backup_path = _backup(excel)
    result["backup_path"] = backup_path
    _log(progress, f"Backed up original to: {backup_path}")

    # 機關全銜 expansion map: built-in abbreviations plus any derived from the
    # master's own 全銜 entries, so new lines carry the full 機關名稱.
    expand_map = build_agency_expansion(existing_prefixes)

    # Only the articles that gained new interpretations get rebuilt + re-sorted;
    # untouched articles are left exactly as they were.
    wrap = Alignment(wrap_text=True, vertical="top")
    for key, items in appended.items():
        if key in existing:
            ridx = existing[key]["row"]
            cell = ws.cell(row=ridx, column=fun_col)
            _assign_text(cell, _merge_sorted_funhao(existing[key]["lines"], items, expand_map))
            cell.alignment = wrap
        else:
            # article_key() hands back the site's raw string when it doesn't
            # look like 第N條. Keep the row (dropping data silently is worse),
            # but say so -- an odd key here means the site changed or the page
            # was tampered with.
            if not _ARTICLE_KEY_RE.fullmatch(key):
                _log(progress, f"  ! 條號格式不符預期，已以純文字寫入：{key[:60]}")
            ridx = ws.max_row + 1
            art_cell = ws.cell(row=ridx, column=art_col)
            _assign_text(art_cell, f"{ARTICLE_LABEL_PREFIX}{key}")
            art_cell.alignment = wrap
            fun_cell = ws.cell(row=ridx, column=fun_col)
            _assign_text(fun_cell, _merge_sorted_funhao([], items, expand_map))
            fun_cell.alignment = wrap

    # Refresh the 最新函釋(共N條) header to the live distinct count.
    funhao_total = count_distinct_funhao(ws, fun_col)
    update_funhao_header(ws, fun_col, funhao_total)
    result["funhao_total"] = funhao_total

    # Atomic write. openpyxl's save() opens the destination with ZipFile(path,
    # 'w'), which TRUNCATES it first -- anything that interrupts the save (the
    # window closed mid-update, antivirus or a sync client grabbing the file)
    # leaves the master sheet as an unopenable stub. Writing a sibling temp file
    # and swapping it in with os.replace() -- atomic on Windows and POSIX alike,
    # and same-directory so it never crosses a filesystem -- means any failure
    # leaves the original exactly as it was.
    tmp_path = f"{excel}.tmp-save"
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, excel)
    except Exception as e:  # noqa: BLE001
        try:
            os.remove(tmp_path)
        except OSError:
            pass
        raise RuntimeError(
            f"寫回失敗，原檔未被更動（備份：{backup_path}）。\n原始錯誤：{e}") from e
    result["written"] = True
    _log(progress, f"Wrote {len(new_rows)} new interpretation(s) to: {excel}; "
                   f"header updated to 共{funhao_total}條.")
    return result


_FUNHAO_COUNT_RE = re.compile(r"\s*[(（]\s*共\s*\d+\s*條\s*[)）]\s*$")


def count_distinct_funhao(ws, fun_col):
    """Total number of DISTINCT 函釋 across the whole 最新函釋 column: a 字號 cited
    under several articles counts once (不重複計算). Uses the same funhao_key
    identity used everywhere else for de-duplication."""
    keys = set()
    for row in ws.iter_rows(min_row=2, values_only=False):
        value = row[fun_col - 1].value
        if not value:
            continue
        for ln in str(value).split("\n"):
            if ln.strip():
                keys.add(funhao_key(ln))
    return len(keys)


def update_funhao_header(ws, fun_col, count):
    """Rewrite the 最新函釋 header cell's '(共N條)' suffix to the live count,
    preserving the base label. Adds the suffix if it was missing."""
    cell = ws.cell(row=1, column=fun_col)
    base = _FUNHAO_COUNT_RE.sub("", str(cell.value or "")).strip()
    cell.value = f"{base}(共{count}條)"


def _funhao_cell_value(items):
    """Build a 函號 cell value from (line, is_new) entries. Newly added lines are
    rendered in red rich text so they stand out at a glance; if nothing is new
    (shouldn't happen here, but kept robust), a plain newline-joined string is
    returned so the cell stays a simple string.

    The line-break is appended to each line's own run rather than emitted as a
    standalone "\\n" run: openpyxl only marks a run xml:space="preserve" when it
    has non-whitespace content, so a whitespace-only run loses its newline and
    Excel then rejects the file as corrupt. Keeping every run non-empty avoids
    that."""
    if not any(is_new for _, is_new in items):
        return "\n".join(text for text, _ in items)
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    red = InlineFont(color=NEW_LINE_COLOR)
    last = len(items) - 1
    parts = []
    for i, (text, is_new) in enumerate(items):
        run = text if i == last else text + "\n"   # newline rides with its line
        parts.append(TextBlock(red, run) if is_new else run)
    return CellRichText(parts)


def _merge_sorted_funhao(existing_lines, new_items, expand_map=None):
    """
    Build a 函號 cell value from existing raw lines plus newly scraped rows:
      - de-duplicate by 字號序號 so a line already present is kept once (its date
        is rewritten to zero-stripped dotted form),
      - expand any abbreviated 機關字軌 on NEWLY added lines to its full 全銜
        (built-in table plus expand_map derived from the master's own entries);
        existing lines are left exactly as they were (簡稱 not touched),
      - sort by 發文日期 ascending so the newest interpretation ends up at the
        bottom (undated lines sort first / oldest, order otherwise preserved).
    Newly added lines are returned in red (see _funhao_cell_value).
    """
    combined = []          # list of (line, is_new)
    seen = set()
    for ln in existing_lines:
        key = funhao_key(ln)
        if key in seen:
            continue
        seen.add(key)
        combined.append((normalize_funhao_line(ln), False))
    for r in new_items:
        key = funhao_key(funhao_line(r))
        if key in seen:
            continue
        seen.add(key)
        combined.append((expand_agency_prefix(funhao_line(r), expand_map), True))
    combined.sort(key=lambda it: parse_roc_date(it[0]) or (-1, -1, -1))
    return _funhao_cell_value(combined)


def _backup(path):
    """Copy the original into a 「備份」subfolder beside it, named
    {name}_backup_{timestamp}{ext}; return the backup path. The timestamp keeps
    the date and time so repeated runs on the same day don't overwrite earlier
    backups. The subfolder is created if it doesn't exist."""
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    folder, filename = os.path.split(path)
    name, ext = os.path.splitext(filename)
    backup_dir = os.path.join(folder, "備份")
    os.makedirs(backup_dir, exist_ok=True)
    backup_path = os.path.join(backup_dir, f"{name}_backup_{stamp}{ext or '.xlsx'}")
    shutil.copy2(path, backup_path)
    return backup_path


# ---------------------------------------------------------------------------
# Dev entry point: quick manual check
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    data = scrape_all(progress=print)
    print(f"\nTotal {len(data)} rows; first 5:")
    for r in data[:5]:
        print(f"  [{r['article']}] {r['doc_no']} | {r['doc_date']}")
